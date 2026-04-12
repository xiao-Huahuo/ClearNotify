from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from app.services.ocr_pipeline_service import extract_text_from_image_bytes_with_fallback

try:
    import docx
except ImportError:
    docx = None


def _load_openpyxl():
    try:
        import openpyxl as _openpyxl

        return _openpyxl, None
    except Exception as exc:
        return None, exc


def _load_xlrd():
    try:
        import xlrd as _xlrd

        return _xlrd, None
    except Exception as exc:
        return None, exc


def _local_name(tag: str) -> str:
    return tag.split("}", 1)[-1] if "}" in tag else tag


def _clean_block_text(text: str) -> str:
    if not text:
        return ""
    normalized = (
        text.replace("\\r\\n", "\n")
        .replace("\\n", "\n")
        .replace("\\r", "\n")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
    )
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in normalized.split("\n")]
    return "\n".join(line for line in lines if line)


def _resolve_related_part(related_parts: dict[str, Any], rel_id: str) -> Any | None:
    if not rel_id:
        return None
    return related_parts.get(rel_id)


def _ocr_docx_image(blob: bytes, *, content_type: str, ocr_cache: dict[str, str]) -> str:
    digest = hashlib.sha1(blob).hexdigest()
    if digest in ocr_cache:
        return ocr_cache[digest]

    result = _clean_block_text(
        extract_text_from_image_bytes_with_fallback(
            blob,
            content_type=content_type or "image/png",
        )
    )
    ocr_cache[digest] = result
    return result


def _extract_image_segments(node, related_parts: dict[str, Any], ocr_cache: dict[str, str]) -> list[str]:
    segments: list[str] = []
    seen_rel_ids: set[str] = set()

    for element in node.iter():
        for attr_name, attr_value in element.attrib.items():
            if not isinstance(attr_value, str) or not attr_value.startswith("rId"):
                continue
            if _local_name(attr_name) not in {"embed", "id", "link"}:
                continue
            if attr_value in seen_rel_ids:
                continue
            seen_rel_ids.add(attr_value)
            part = _resolve_related_part(related_parts, attr_value)
            blob = getattr(part, "blob", None)
            if not blob:
                continue
            text = _ocr_docx_image(
                blob,
                content_type=str(getattr(part, "content_type", "") or "image/png"),
                ocr_cache=ocr_cache,
            )
            if text:
                segments.append(text)
    return segments


def _append_clean_text(buffer: list[str], segments: list[str]) -> None:
    text = _clean_block_text("".join(buffer))
    buffer.clear()
    if text:
        segments.append(text)


def _extract_paragraph_segments(paragraph, related_parts: dict[str, Any], ocr_cache: dict[str, str]) -> list[str]:
    segments: list[str] = []
    buffer: list[str] = []

    def walk(node) -> None:
        node_name = _local_name(node.tag)
        if node_name == "r":
            for child in node.iterchildren():
                child_name = _local_name(child.tag)
                if child_name == "t":
                    buffer.append(child.text or "")
                elif child_name == "tab":
                    buffer.append("\t")
                elif child_name in {"br", "cr"}:
                    buffer.append("\n")
                elif child_name in {"drawing", "pict"}:
                    _append_clean_text(buffer, segments)
                    segments.extend(_extract_image_segments(child, related_parts, ocr_cache))
            return

        if node_name == "hyperlink":
            for child in node.iterchildren():
                walk(child)
            return

        for child in node.iterchildren():
            walk(child)

    for child in paragraph.iterchildren():
        walk(child)

    _append_clean_text(buffer, segments)
    return segments


def _extract_table_text(table, related_parts: dict[str, Any], ocr_cache: dict[str, str]) -> str:
    rows: list[str] = []

    for row in table.iterchildren():
        if _local_name(row.tag) != "tr":
            continue
        cells: list[str] = []
        for cell in row.iterchildren():
            if _local_name(cell.tag) != "tc":
                continue
            blocks = _extract_docx_blocks(cell, related_parts, ocr_cache)
            cell_text = " / ".join(blocks).strip()
            cells.append(cell_text)
        if any(cell.strip() for cell in cells):
            rows.append(" | ".join(cell.strip() for cell in cells))
    return "\n".join(rows).strip()


def _extract_docx_blocks(container, related_parts: dict[str, Any], ocr_cache: dict[str, str]) -> list[str]:
    blocks: list[str] = []

    for child in container.iterchildren():
        child_name = _local_name(child.tag)
        if child_name == "p":
            blocks.extend(_extract_paragraph_segments(child, related_parts, ocr_cache))
            continue
        if child_name == "tbl":
            table_text = _extract_table_text(child, related_parts, ocr_cache)
            if table_text:
                blocks.append(table_text)
    return [block for block in blocks if block.strip()]


def count_docx_inline_images(file_path: Path) -> int:
    try:
        with ZipFile(file_path) as archive:
            return sum(
                1
                for name in archive.namelist()
                if name.startswith("word/media/") and not name.endswith("/")
            )
    except Exception:
        return 0


def extract_text_from_docx(file_path: Path) -> str:
    if docx is None:
        return "python-docx is not installed, so DOCX content cannot be extracted."

    try:
        document = docx.Document(file_path)
        related_parts = getattr(document.part, "related_parts", {})
        ocr_cache: dict[str, str] = {}
        blocks = _extract_docx_blocks(document.element.body, related_parts, ocr_cache)
        return "\n".join(blocks)
    except Exception as exc:
        return f"DOCX parse failed: {exc}"


def extract_text_from_excel(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    full_text: list[str] = []

    try:
        if ext == ".xlsx":
            openpyxl, openpyxl_err = _load_openpyxl()
            if openpyxl is None:
                return f"Excel parse failed: unable to load openpyxl ({openpyxl_err})"

            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                full_text.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_data = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                    if row_data:
                        full_text.append(" | ".join(row_data))

        elif ext == ".xls":
            xlrd, xlrd_err = _load_xlrd()
            if xlrd is None:
                return f"Excel parse failed: unable to load xlrd ({xlrd_err})"

            wb = xlrd.open_workbook(file_path)
            for sheet in wb.sheets():
                full_text.append(f"--- Sheet: {sheet.name} ---")
                for row_idx in range(sheet.nrows):
                    row_values = sheet.row_values(row_idx)
                    row_data = [str(val).strip() for val in row_values if val and str(val).strip()]
                    if row_data:
                        full_text.append(" | ".join(row_data))
        else:
            return f"Excel parse failed: unsupported suffix {ext}"

        return "\n".join(full_text)
    except Exception as exc:
        return f"Excel parse failed: {exc}"
